"""Tests for KV1: Zusatzfragen (registration_questions / extra_answers)."""

from __future__ import annotations

from uuid import uuid4

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.repositories import organizations
from app.repositories import registrations as registrations_repo
from app.repositories.registrations import RegistrationTarget
from app.repositories import camps as camps_repo

from .auth_helpers import owner_headers as _owner_headers, org_admin_headers as _org_admin_headers


ORG_SLUG = "test-fc"
CAMP_SLUG = "sommer-1"


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------

def _org_row(slug: str = ORG_SLUG, registration_questions: list | None = None) -> dict:
    return {
        "id": uuid4(),
        "slug": slug,
        "name": "Test FC",
        "legal_name": None,
        "legal_address": None,
        "contact_email": "test@example.com",
        "contact_phone": None,
        "logo_url": None,
        "primary_color": None,
        "plan_status": "active", "review_status": "approved",
        "site_published": True,
        "registration_questions": registration_questions if registration_questions is not None else [],
    }


def _target(**overrides) -> RegistrationTarget:
    from datetime import date
    base = dict(
        id=uuid4(),
        organization_id=uuid4(),
        slug=CAMP_SLUG,
        start_date=date(2027, 7, 5),
        end_date=date(2027, 7, 9),
        registration_start=None,
        registration_end=None,
        age_min=6,
        age_max=12,
        capacity=20,
        status="published",
    )
    base.update(overrides)
    return RegistrationTarget(**base)


def _valid_payload(**overrides) -> dict:
    base = {
        "parent_first_name": "Max",
        "parent_last_name": "Mustermann",
        "parent_email": "max@example.com",
        "parent_phone": "+49 123 456789",
        "child_first_name": "Lena",
        "child_last_name": "Mustermann",
        "child_birth_date": "2017-05-10",
        "emergency_contact_phone": "+49 987 654321",
        "photo_permission": False,
        "terms_accepted": True,
        "privacy_accepted": True,
    }
    base.update(overrides)
    return base


def _fake_row(**overrides) -> dict:
    base = {
        "registration_token": uuid4(),
        "status": "registered",
        "payment_status": "open",
        "child_first_name": "Lena",
        "child_last_name": "Mustermann",
    }
    base.update(overrides)
    return base


def _post(client: TestClient, payload: dict):
    return client.post(
        f"/api/v1/organizations/{ORG_SLUG}/camps/{CAMP_SLUG}/registrations",
        json=payload,
    )


# ---------------------------------------------------------------------------
# 1. Schema-Validierung für registration_questions
# ---------------------------------------------------------------------------

class TestQuestionSchemaValidation:
    """
    Punkt 2: Fragetypen validiert in Pydantic.
    Punkt 8: '6. Frage → 422, ungültiger Typ → 422'.
    """

    def test_six_questions_rejected_with_422(self):
        """Mehr als 5 Fragen → 422 beim Speichern der Vereinskonfiguration."""
        headers = _owner_headers()
        six_questions = [
            {"key": f"q{i}", "label": f"Frage {i}", "type": "text", "required": False}
            for i in range(6)
        ]
        with TestClient(app) as client:
            response = client.patch(
                f"/admin/organizations/{ORG_SLUG}",
                json={"registration_questions": six_questions},
                headers=headers,
            )
        assert response.status_code == 422

    def test_invalid_type_rejected_with_422(self):
        """Ungültiger Fragetyp → 422."""
        headers = _owner_headers()
        with TestClient(app) as client:
            response = client.patch(
                f"/admin/organizations/{ORG_SLUG}",
                json={"registration_questions": [
                    {"key": "q1", "label": "Foo", "type": "invalid_type", "required": True}
                ]},
                headers=headers,
            )
        assert response.status_code == 422

    def test_five_questions_accepted(self, monkeypatch):
        """Genau 5 Fragen sind erlaubt (Pydantic-Validierung passiert)."""
        five_questions = [
            {"key": f"q{i}", "label": f"Frage {i}", "type": "text", "required": False}
            for i in range(5)
        ]
        monkeypatch.setattr(organizations, "update_organization", lambda slug, data: {
            **_org_row(slug),
            "billing_notes": None,
            "site_published": False,
            "theme": "tradition",
            "iban": None, "intro_heading": None, "intro_text": None,
            "hero_image_url": None, "trainer_image_url": None, "trainer_text": None,
            "terms_text": None, "contact_person_name": None,
            "registration_questions": five_questions,
        })
        headers = _owner_headers()
        with TestClient(app) as client:
            response = client.patch(
                f"/admin/organizations/{ORG_SLUG}",
                json={"registration_questions": five_questions},
                headers=headers,
            )
        # Pydantic validation passes (5 is OK); DB call returns 200
        assert response.status_code != 422

    def test_choice_without_options_rejected(self):
        """Choice-Frage ohne Optionen → 422."""
        headers = _owner_headers()
        with TestClient(app) as client:
            response = client.patch(
                f"/admin/organizations/{ORG_SLUG}",
                json={"registration_questions": [
                    {"key": "q1", "label": "Foo", "type": "choice", "options": [], "required": True}
                ]},
                headers=headers,
            )
        assert response.status_code == 422

    def test_invalid_key_format_rejected(self):
        """Schlüssel mit Großbuchstaben → 422."""
        headers = _owner_headers()
        with TestClient(app) as client:
            response = client.patch(
                f"/admin/organizations/{ORG_SLUG}",
                json={"registration_questions": [
                    {"key": "BadKey", "label": "Foo", "type": "text", "required": False}
                ]},
                headers=headers,
            )
        assert response.status_code == 422


# ---------------------------------------------------------------------------
# 2. Anmeldung mit extra_answers - Pflichtprüfung
# ---------------------------------------------------------------------------

KSV_QUESTIONS = [
    {"key": "position", "label": "Position", "type": "choice",
     "options": ["Feldspieler", "Torwart"], "other_label": None, "required": True},
    {"key": "verein", "label": "Verein", "type": "choice",
     "options": ["KSV Baunatal"], "other_label": "Anderer Verein", "required": True},
]


class TestExtraAnswersValidation:
    """Punkt 8: Anmeldung ohne Pflichtantwort → 422, unbekannter key → 422,
    Sonstiges ohne Text → 422. Plus Regression: Verein ohne Fragen = OK."""

    def test_missing_required_answer_raises_error(self):
        """Pflichtfrage ohne Antwort → ExtraAnswersInvalidError (Mutationsprobe)."""
        from app.repositories.registrations import validate_extra_answers, ExtraAnswersInvalidError
        with pytest.raises(ExtraAnswersInvalidError):
            validate_extra_answers(KSV_QUESTIONS, {})

    def test_unknown_key_raises_error(self):
        """Unbekannter Key in extra_answers → ExtraAnswersInvalidError."""
        from app.repositories.registrations import validate_extra_answers, ExtraAnswersInvalidError
        with pytest.raises(ExtraAnswersInvalidError, match="Unknown answer key"):
            validate_extra_answers(
                KSV_QUESTIONS,
                {"position": "Feldspieler", "verein": "KSV Baunatal", "unknown_field": "foo"},
            )

    def test_invalid_choice_rejected(self):
        """Ungültige Choice-Antwort (kein other_label) → ExtraAnswersInvalidError."""
        from app.repositories.registrations import validate_extra_answers, ExtraAnswersInvalidError
        with pytest.raises(ExtraAnswersInvalidError):
            validate_extra_answers(
                [{"key": "pos", "label": "Pos", "type": "choice",
                  "options": ["Feldspieler", "Torwart"], "other_label": None, "required": True}],
                {"pos": "Schiedsrichter"},
            )

    def test_sonstiges_without_text_rejected_when_required(self):
        """Sonstiges ohne Text → 422 (Antwort = leer, Pflichtfeld)."""
        from app.repositories.registrations import validate_extra_answers, ExtraAnswersInvalidError
        # When other_label is set but user enters empty string
        with pytest.raises(ExtraAnswersInvalidError):
            validate_extra_answers(
                [{"key": "verein", "label": "Verein", "type": "choice",
                  "options": ["KSV Baunatal"], "other_label": "Anderer Verein", "required": True}],
                {"verein": ""},  # empty = not answered, but required
            )

    def test_sonstiges_with_text_accepted(self):
        """Sonstiges mit Text → valide."""
        from app.repositories.registrations import validate_extra_answers
        # Should not raise
        validate_extra_answers(
            [{"key": "verein", "label": "Verein", "type": "choice",
              "options": ["KSV Baunatal"], "other_label": "Anderer Verein", "required": True}],
            {"verein": "FC Kassel"},  # not in options, but other_label set → accepted
        )

    def test_valid_answers_accepted(self):
        """Alle Antworten korrekt → kein Fehler."""
        from app.repositories.registrations import validate_extra_answers
        validate_extra_answers(
            KSV_QUESTIONS,
            {"position": "Feldspieler", "verein": "KSV Baunatal"},
        )

    def test_org_without_questions_works_as_before(self, monkeypatch):
        """Verein ohne Fragen: Anmeldung ohne extra_answers → 201 (Regression)."""
        monkeypatch.setattr(
            organizations, "get_organization_by_slug",
            lambda slug: _org_row(slug, registration_questions=[]),
        )
        monkeypatch.setattr(registrations_repo, "get_registration_target", lambda t, s: _target())
        monkeypatch.setattr(
            registrations_repo, "create_registration",
            lambda tenant, camp, data: _fake_row(),
        )

        with TestClient(app) as client:
            response = _post(client, _valid_payload())
        assert response.status_code == 201

    def test_payment_reference_excludes_child_name(self, monkeypatch):
        """The selected camp and CP reference identify payment without child names."""
        monkeypatch.setattr(
            organizations, "get_organization_by_slug",
            lambda slug: _org_row(slug, registration_questions=[]),
        )
        monkeypatch.setattr(
            registrations_repo, "get_registration_target",
            lambda t, s: _target(title="Winter-Fußballschule 2027"),
        )
        monkeypatch.setattr(
            registrations_repo, "create_registration",
            lambda tenant, camp, data: _fake_row(
                child_first_name=data.child_first_name,
                child_last_name=data.child_last_name,
            ),
        )

        with TestClient(app) as client:
            response = _post(client, _valid_payload(child_first_name="Lena", child_last_name="Muster"))

        assert response.status_code == 201
        body = response.json()
        ref = body["payment_reference"]
        assert ref == f"Winter-Fußballschule 2027 CP-{body['registration_token'][:8].upper()}"
        assert "Lena" not in ref
        assert "Muster" not in ref
        assert len(ref) <= 140

    def test_payment_reference_excludes_child_name_fields_from_response(self, monkeypatch):
        """child_first_name und child_last_name sind nicht im API-Response (exclude=True)."""
        monkeypatch.setattr(
            organizations, "get_organization_by_slug",
            lambda slug: _org_row(slug, registration_questions=[]),
        )
        monkeypatch.setattr(registrations_repo, "get_registration_target", lambda t, s: _target())
        monkeypatch.setattr(
            registrations_repo, "create_registration",
            lambda tenant, camp, data: _fake_row(
                child_first_name=data.child_first_name,
                child_last_name=data.child_last_name,
            ),
        )

        with TestClient(app) as client:
            response = _post(client, _valid_payload())

        assert response.status_code == 201
        body = response.json()
        assert set(body.keys()) == {"registration_token", "status", "payment_status", "payment_reference"}

    def test_text_answer_too_long_is_trimmed_by_schema(self, monkeypatch):
        """Text-Antwort > 80 Zeichen wird durch RegistrationCreate._trim_and_limit_answers
        auf 80 Zeichen gekürzt (kein Fehler, nur Trimmung)."""
        from app.schemas import RegistrationCreate
        import pydantic

        data = RegistrationCreate(
            parent_first_name="Max",
            parent_last_name="Muster",
            parent_email="max@example.com",
            parent_phone="+49123",
            child_first_name="Lena",
            child_last_name="Muster",
            child_birth_date="2017-01-01",
            emergency_contact_phone="+49987",
            photo_permission=False,
            terms_accepted=True,
            privacy_accepted=True,
            extra_answers={"key": "x" * 100},
        )
        assert len(data.extra_answers["key"]) == 80


# ---------------------------------------------------------------------------
# 3. Export enthält extra_answers als Spalten
# ---------------------------------------------------------------------------

class TestExportWithExtraAnswers:
    """Punkt 8: Antworten im Export als Spalten."""

    def test_participants_export_contains_extra_answer_columns(self, monkeypatch):
        """Extra-Antwort-Spalten erscheinen im Teilnehmerliste-Export."""
        from io import BytesIO
        from openpyxl import load_workbook
        from datetime import date, datetime, timezone

        org_id = uuid4()
        camp_id = uuid4()
        org = {
            **_org_row(), "id": org_id, "site_published": True,
            "billing_notes": None, "theme": "tradition", "iban": None,
            "intro_heading": None, "intro_text": None, "hero_image_url": None,
            "trainer_image_url": None, "trainer_text": None, "terms_text": None,
            "contact_person_name": None,
            "registration_questions": [
                {"key": "position", "label": "Position", "type": "choice",
                 "options": ["Feldspieler", "Torwart"], "other_label": None, "required": True}
            ],
        }

        monkeypatch.setattr(organizations, "get_organization_by_slug", lambda slug: org)
        monkeypatch.setattr(camps_repo, "get_camp_by_slug", lambda org_id, slug: {
            "id": camp_id, "organization_id": org_id, "slug": CAMP_SLUG,
            "title": "Sommer", "start_date": "2027-07-05", "end_date": "2027-07-09",
            "price_cents": 5000, "currency": "EUR", "status": "published",
        })
        monkeypatch.setattr(registrations_repo, "list_registrations_for_camp", lambda org_id, c_id: [{
            "id": uuid4(), "registration_token": uuid4(),
            "status": "registered", "payment_status": "open",
            "parent_first_name": "Max", "parent_last_name": "Muster",
            "parent_email": "max@x.de", "parent_phone": "+49123",
            "child_first_name": "Lena", "child_last_name": "Muster",
            "child_birth_date": date(2017, 1, 1),
            "emergency_contact_name": None, "emergency_contact_phone": None,
            "medical_notes": None, "allergies": None, "jersey_size": None, "pickup_authorized": None,
            "photo_permission": False,
            "created_at": datetime(2027, 1, 1, tzinfo=timezone.utc),
            "extra_answers": {"position": "Torwart"},
        }])

        headers = _owner_headers()
        with TestClient(app) as client:
            response = client.get(
                f"/admin/organizations/{ORG_SLUG}/camps/{CAMP_SLUG}/export.xlsx?view=participants",
                headers=headers,
            )

        assert response.status_code == 200
        wb = load_workbook(BytesIO(response.content))
        ws = wb.active
        header_row = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
        assert "Position" in header_row
        # Find the column
        pos_col = header_row.index("Position") + 1
        assert ws.cell(2, pos_col).value == "Torwart"

    def test_daylist_export_contains_extra_answer_columns(self, monkeypatch):
        """Extra-Antwort-Spalten erscheinen auch im Tagesliste-Export."""
        from io import BytesIO
        from openpyxl import load_workbook
        from datetime import date, datetime, timezone

        org_id = uuid4()
        camp_id = uuid4()
        org = {
            **_org_row(), "id": org_id, "site_published": True,
            "billing_notes": None, "theme": "tradition", "iban": None,
            "intro_heading": None, "intro_text": None, "hero_image_url": None,
            "trainer_image_url": None, "trainer_text": None, "terms_text": None,
            "contact_person_name": None,
            "registration_questions": [
                {"key": "position", "label": "Position", "type": "choice",
                 "options": ["Feldspieler", "Torwart"], "other_label": None, "required": True}
            ],
        }

        monkeypatch.setattr(organizations, "get_organization_by_slug", lambda slug: org)
        monkeypatch.setattr(camps_repo, "get_camp_by_slug", lambda org_id, slug: {
            "id": camp_id, "organization_id": org_id, "slug": CAMP_SLUG,
            "title": "Sommer", "start_date": "2027-07-05", "end_date": "2027-07-09",
            "price_cents": 5000, "currency": "EUR", "status": "published",
        })
        monkeypatch.setattr(registrations_repo, "list_registrations_for_camp", lambda org_id, c_id: [{
            "id": uuid4(), "registration_token": uuid4(),
            "status": "registered", "payment_status": "open",
            "parent_first_name": "Max", "parent_last_name": "Muster",
            "parent_email": "max@x.de", "parent_phone": "+49123",
            "child_first_name": "Lena", "child_last_name": "Muster",
            "child_birth_date": date(2017, 1, 1),
            "emergency_contact_name": None, "emergency_contact_phone": None,
            "medical_notes": None, "allergies": None, "jersey_size": None, "pickup_authorized": None,
            "photo_permission": False,
            "created_at": datetime(2027, 1, 1, tzinfo=timezone.utc),
            "extra_answers": {"position": "Feldspieler"},
        }])

        headers = _owner_headers()
        with TestClient(app) as client:
            response = client.get(
                f"/admin/organizations/{ORG_SLUG}/camps/{CAMP_SLUG}/export.xlsx?view=daylist",
                headers=headers,
            )

        assert response.status_code == 200
        wb = load_workbook(BytesIO(response.content))
        ws = wb.active
        header_row = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
        assert "Position" in header_row
        pos_col = header_row.index("Position") + 1
        assert ws.cell(2, pos_col).value == "Feldspieler"


# ---------------------------------------------------------------------------
# 4. Tenant-Isolation: org_admin von Verein A kann Fragen von Verein B nicht ändern
# ---------------------------------------------------------------------------

class TestTenantIsolation:
    """Punkt 8: org_admin von Verein A kann Fragen von Verein B nicht ändern (403)."""

    def test_org_admin_cannot_update_other_orgs_questions(self, monkeypatch):
        """Org-Admin von 'test-fc' kann Fragen von 'other-fc' nicht ändern."""
        other_org_id = uuid4()
        my_org_id = uuid4()
        other_org = {
            **_org_row("other-fc"), "id": other_org_id, "site_published": True,
            "billing_notes": None, "theme": "tradition", "iban": None,
            "intro_heading": None, "intro_text": None, "hero_image_url": None,
            "trainer_image_url": None, "trainer_text": None, "terms_text": None,
            "contact_person_name": None,
        }
        monkeypatch.setattr(organizations, "get_organization_by_slug", lambda slug: other_org)

        # org_admin for my_org_id, NOT other_org_id
        headers = _org_admin_headers(str(my_org_id))

        with TestClient(app) as client:
            response = client.patch(
                "/admin/organizations/other-fc",
                json={"registration_questions": []},
                headers=headers,
            )
        assert response.status_code == 403
