from __future__ import annotations

from datetime import date, datetime, timezone
from io import BytesIO
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.repositories import camps as camps_repo
from app.repositories import organizations
from app.repositories import registrations as registrations_repo

from .auth_helpers import owner_headers as _auth_headers


def _org_row(slug: str = "demo-fc") -> dict:
    return {
        "id": uuid4(),
        "slug": slug,
        "name": "Demo Football Academy",
        "legal_name": None,
        "contact_email": "demo@example.com",
        "contact_phone": None,
        "logo_url": None,
        "primary_color": None,
        "plan_status": "pilot",
    }


def _camp_row(organization_id, slug: str = "summer-1", title: str = "Summer Camp Week 1", price_cents: int = 5000) -> dict:
    return {
        "id": uuid4(),
        "organization_id": organization_id,
        "slug": slug,
        "title": title,
        "start_date": "2027-07-05",
        "end_date": "2027-07-09",
        "registration_start": None,
        "registration_end": None,
        "age_min": 6,
        "age_max": 12,
        "capacity": 20,
        "price_cents": price_cents,
        "currency": "EUR",
        "status": "published",
    }


def _registration_row(**overrides) -> dict:
    base = {
        "id": uuid4(),
        "registration_token": uuid4(),
        "status": "registered",
        "payment_status": "open",
        "parent_first_name": "Max",
        "parent_last_name": "Mustermann",
        "parent_email": "max@example.com",
        "parent_phone": "+49 123 456789",
        "child_first_name": "Lena",
        "child_last_name": "Mustermann",
        "child_birth_date": date(2018, 5, 10),
        "emergency_contact_name": None,
        "emergency_contact_phone": None,
        "medical_notes": None,
        "allergies": None,
        "jersey_size": None,
        "pickup_authorized": None,
        "photo_permission": False,
        "created_at": datetime(2027, 1, 1, tzinfo=timezone.utc),
    }
    base.update(overrides)
    return base


def _load(content: bytes):
    return load_workbook(BytesIO(content))


def test_export_camp_participants_xlsx(monkeypatch):
    org = _org_row()
    camp = _camp_row(org["id"])
    monkeypatch.setattr(organizations, "get_organization_by_slug", lambda slug: org)
    monkeypatch.setattr(camps_repo, "get_camp_by_slug", lambda org_id, slug: camp)
    monkeypatch.setattr(
        registrations_repo,
        "list_registrations_for_camp",
        lambda org_id, camp_id: [_registration_row(status="registered"), _registration_row(status="waitlist")],
    )
    headers = _auth_headers()

    with TestClient(app) as client:
        response = client.get("/admin/organizations/demo-fc/camps/summer-1/export.xlsx", headers=headers)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "summer-1-participants.xlsx" in response.headers["content-disposition"]
    wb = _load(response.content)
    sheet = wb.active
    assert sheet["A1"].value == "Status"
    assert sheet.max_row == 3  # header + 2 registrations


def test_export_camp_payments_xlsx_excludes_camp_column(monkeypatch):
    org = _org_row()
    camp = _camp_row(org["id"], price_cents=6000)
    monkeypatch.setattr(organizations, "get_organization_by_slug", lambda slug: org)
    monkeypatch.setattr(camps_repo, "get_camp_by_slug", lambda org_id, slug: camp)
    monkeypatch.setattr(
        registrations_repo,
        "list_registrations_for_camp",
        lambda org_id, camp_id: [_registration_row(payment_status="paid")],
    )
    headers = _auth_headers()

    with TestClient(app) as client:
        response = client.get("/admin/organizations/demo-fc/camps/summer-1/export.xlsx?view=payments", headers=headers)

    assert response.status_code == 200
    sheet = _load(response.content).active
    assert [c.value for c in sheet[1]] == ["Kind", "Elternteil", "E-Mail", "Preis (EUR)", "Zahlungsstatus"]
    assert sheet["D2"].value == 60.0
    assert sheet["E2"].value == "Bezahlt"


def test_export_camp_waitlist_xlsx_only_includes_waitlisted(monkeypatch):
    org = _org_row()
    camp = _camp_row(org["id"])
    monkeypatch.setattr(organizations, "get_organization_by_slug", lambda slug: org)
    monkeypatch.setattr(camps_repo, "get_camp_by_slug", lambda org_id, slug: camp)
    monkeypatch.setattr(
        registrations_repo,
        "list_registrations_for_camp",
        lambda org_id, camp_id: [
            _registration_row(status="registered", child_first_name="Registered"),
            _registration_row(status="waitlist", child_first_name="Waiting"),
        ],
    )
    headers = _auth_headers()

    with TestClient(app) as client:
        response = client.get("/admin/organizations/demo-fc/camps/summer-1/export.xlsx?view=waitlist", headers=headers)

    assert response.status_code == 200
    sheet = _load(response.content).active
    assert sheet.max_row == 2  # header + exactly one waitlisted row
    assert "Waiting" in sheet["B2"].value


def test_export_camp_xlsx_unknown_camp_returns_404(monkeypatch):
    org = _org_row()
    monkeypatch.setattr(organizations, "get_organization_by_slug", lambda slug: org)
    monkeypatch.setattr(camps_repo, "get_camp_by_slug", lambda org_id, slug: None)
    headers = _auth_headers()

    with TestClient(app) as client:
        response = client.get("/admin/organizations/demo-fc/camps/does-not-exist/export.xlsx", headers=headers)

    assert response.status_code == 404


def test_export_camp_xlsx_without_auth_returns_401():
    with TestClient(app) as client:
        response = client.get("/admin/organizations/demo-fc/camps/summer-1/export.xlsx")

    assert response.status_code == 401


def test_export_organization_payments_xlsx_includes_camp_column(monkeypatch):
    org = _org_row()
    camp_a = _camp_row(org["id"], slug="week-1", title="Week 1", price_cents=5000)
    camp_b = _camp_row(org["id"], slug="week-2", title="Week 2", price_cents=7000)
    monkeypatch.setattr(organizations, "get_organization_by_slug", lambda slug: org)
    monkeypatch.setattr(camps_repo, "list_camps_for_organization", lambda org_id: [camp_a, camp_b])
    monkeypatch.setattr(
        registrations_repo,
        "list_registrations_for_organization",
        lambda org_id: [
            _registration_row(camp_id=camp_a["id"]),
            _registration_row(camp_id=camp_b["id"], payment_status="paid"),
        ],
    )
    headers = _auth_headers()

    with TestClient(app) as client:
        response = client.get("/admin/organizations/demo-fc/export.xlsx?view=payments", headers=headers)

    assert response.status_code == 200
    sheet = _load(response.content).active
    assert sheet[1][0].value == "Camp"
    assert sheet.max_row == 3  # header + one row per camp
    camp_names = {sheet.cell(row=r, column=1).value for r in (2, 3)}
    assert camp_names == {"Week 1", "Week 2"}


def test_export_organization_waitlist_xlsx_without_auth_returns_401():
    with TestClient(app) as client:
        response = client.get("/admin/organizations/demo-fc/export.xlsx?view=waitlist")

    assert response.status_code == 401
