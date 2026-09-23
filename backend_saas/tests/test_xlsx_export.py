from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.xlsx_export import build_xlsx


def test_build_xlsx_writes_header_row():
    content = build_xlsx(["Name", "Alter"], [["Lena", 8], ["Max", 10]], sheet_name="Teilnehmer")
    wb = load_workbook(BytesIO(content))
    sheet = wb.active

    assert sheet.title == "Teilnehmer"
    assert [c.value for c in sheet[1]] == ["Name", "Alter"]
    assert [c.value for c in sheet[2]] == ["Lena", 8]
    assert [c.value for c in sheet[3]] == ["Max", 10]


def test_build_xlsx_handles_empty_rows():
    content = build_xlsx(["Name"], [])
    wb = load_workbook(BytesIO(content))
    sheet = wb.active

    assert [c.value for c in sheet[1]] == ["Name"]
    assert sheet.max_row == 1


def test_build_xlsx_truncates_long_sheet_name():
    content = build_xlsx(["A"], [], sheet_name="x" * 50)
    wb = load_workbook(BytesIO(content))
    assert len(wb.active.title) == 31


def test_build_xlsx_freezes_header_row():
    content = build_xlsx(["A"], [["b"]])
    wb = load_workbook(BytesIO(content))
    assert wb.active.freeze_panes == "A2"


def test_build_xlsx_neutralizes_formula_trigger_prefixes():
    """Parent-supplied free text (allergies, notes, names, ...) could start
    with '=', '+', '-', '@', a tab, or a CR — Excel/LibreOffice would
    otherwise evaluate that as a formula when an admin opens the export."""
    content = build_xlsx(
        ["Hinweis"],
        [
            ["=HYPERLINK(\"http://evil\",\"click\")"],
            ["+1+1"],
            ["-1+1"],
            ["@SUM(1,1)"],
            ["harmless text"],
        ],
    )
    wb = load_workbook(BytesIO(content))
    sheet = wb.active
    assert sheet["A2"].value == "'=HYPERLINK(\"http://evil\",\"click\")"
    assert sheet["A3"].value == "'+1+1"
    assert sheet["A4"].value == "'-1+1"
    assert sheet["A5"].value == "'@SUM(1,1)"
    assert sheet["A6"].value == "harmless text"
