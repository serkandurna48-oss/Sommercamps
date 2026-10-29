"""
UC4b Ablauftest: Verein + Camp anlegen → Elternanmeldungen → Admin-Operationen.

Gegen lokale Test-Datenbank (REVIEW_POSTGRES_PORT, z. B. 55443).
Keine echten E-Mails (FakeProvider), kein Stripe, kein echtes Supabase Auth
(supabase_auth.verify_access_token wird für synthetische Identitäten gepacht).

Ablauf:
  1. Verein (billing_notes + IBAN) anlegen, veröffentlichen
  2. Camp (Kapazität 2) anlegen, veröffentlichen
  3. billing_notes nicht in Vereins-Admin-Antwort (Org-Admin-Rolle)
  4. Kind 1 → status=registered, CP-Referenz, Bestätigungsmail (kein Name im Betreff)
  5. Kind 2 → status=registered (Camp jetzt voll)
  6. Kind 3 → status=waitlist, Wartelistenmail
  7. Admin: Kind 1 als bezahlt markieren → payment_status=paid
  8. Admin: Kind 2 stornieren → Kind 3 rückt nach (status=registered), API meldet promoted
  9. Tagesliste enthält Kind 1 + Kind 3, nicht Kind 2
  10. Doppelte Anmeldung (Kind 1 nochmals) → 409
  11. Fremder Vereins-Admin → 403
  12. billing_notes nie in öffentlicher API

Läuft nur wenn REVIEW_POSTGRES_PORT gesetzt ist; sonst skip (0 failures).
"""

from __future__ import annotations

import os
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import psycopg2
from psycopg2 import sql
from psycopg2.extras import RealDictCursor, register_uuid
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import db, supabase_auth
from app.main import app
from app.notifications.mail import FakeProvider, reset_provider_for_testing


@pytest.fixture(scope="module")
def local_database():
    port = os.environ.get("REVIEW_POSTGRES_PORT")
    if not port:
        pytest.skip("Set REVIEW_POSTGRES_PORT to opt into disposable local DB tests")
    assert port.isdigit() and 1024 < int(port) < 65536
    config = dict(host="127.0.0.1", port=int(port), user="postgres", connect_timeout=5)
    name = f"journey_campspilot_{uuid4().hex}"
    admin = psycopg2.connect(dbname="postgres", **config)
    admin.autocommit = True
    register_uuid()
    with admin.cursor() as cur:
        cur.execute(sql.SQL("CREATE DATABASE {}").format(sql.Identifier(name)))
    try:
        with psycopg2.connect(dbname=name, **config) as setup:
            with setup.cursor() as cur:
                cur.execute("CREATE SCHEMA auth; CREATE TABLE auth.users (id uuid PRIMARY KEY)")
                cur.execute(
                    """CREATE FUNCTION auth.uid() RETURNS uuid LANGUAGE sql STABLE
                    AS $$ SELECT (nullif(current_setting('request.jwt.claims', true), '')::jsonb
                        ->> 'sub')::uuid $$"""
                )
                for role in ("anon", "authenticated"):
                    cur.execute("SELECT 1 FROM pg_roles WHERE rolname = %s", (role,))
                    if cur.fetchone() is None:
                        cur.execute(
                            sql.SQL("CREATE ROLE {} NOLOGIN").format(sql.Identifier(role))
                        )
                migrations = Path(__file__).resolve().parents[2] / "supabase" / "migrations"
                for migration in sorted(migrations.glob("*.sql")):
                    cur.execute(migration.read_text(encoding="utf-8"))
                cur.execute("GRANT USAGE ON SCHEMA auth TO anon, authenticated")
                cur.execute("GRANT USAGE ON SCHEMA public TO anon, authenticated")
                cur.execute(
                    "GRANT SELECT, INSERT, UPDATE, DELETE ON ALL TABLES IN SCHEMA public"
                    " TO anon, authenticated"
                )
        setup.close()
        yield dict(dbname=name, **config)
    finally:
        with admin.cursor() as cur:
            cur.execute(
                sql.SQL("DROP DATABASE {} WITH (FORCE)").format(sql.Identifier(name))
            )
        admin.close()


def test_uc4b_end_to_end_journey(local_database, monkeypatch):
    """Vollständiger UC4b-Ablauf gegen echte Postgres-Datenbank.

    Deckt ab: Verein/Camp-Anlage, Anmeldungs-Status (registered/waitlist),
    CP-Referenz, Mail-Betreff ohne PII, Bezahlt-Markierung, Stornierung mit
    automatischem Nachrücken, Tageslisten-Export, Doppel-409,
    Tenant-Isolation (403), billing_notes-Datenleck.
    """

    @contextmanager
    def cursor():
        conn = psycopg2.connect(**local_database)
        try:
            with conn:
                with conn.cursor(cursor_factory=RealDictCursor) as cur:
                    cur.execute("SET LOCAL statement_timeout = '8s'")
                    yield cur
        finally:
            conn.close()

    monkeypatch.setattr(db, "get_cursor", cursor)

    owner_id, admin_a_id, admin_b_id = uuid4(), uuid4(), uuid4()
    identities = {
        "owner": owner_id,
        "admin-a": admin_a_id,
        "admin-b": admin_b_id,
    }

    with cursor() as cur:
        for uid in identities.values():
            cur.execute("INSERT INTO auth.users(id) VALUES (%s)", (uid,))
        cur.execute("INSERT INTO platform_owners(user_id) VALUES (%s)", (owner_id,))

    def verify(token):
        if token not in identities:
            raise supabase_auth.SupabaseAuthError("Unknown synthetic identity")
        return supabase_auth.SupabaseUser(str(identities[token]), f"{token}@example.com")

    monkeypatch.setattr(supabase_auth, "verify_access_token", verify)

    fake = FakeProvider()
    reset_provider_for_testing(fake)
    try:
        _run_journey(cursor, identities, fake)
    finally:
        reset_provider_for_testing(None)


def _run_journey(cursor, identities, fake: FakeProvider) -> None:
    admin_a_id = identities["admin-a"]
    h = lambda token: {"Authorization": f"Bearer {token}"}

    with TestClient(app) as client:
        def req(method, path, expected=200, identity="admin-a", **kwargs):
            resp = client.request(
                method, path,
                headers=h(identity) if identity else {},
                **kwargs,
            )
            assert resp.status_code == expected, (
                f"{method} {path} → {resp.status_code}: {resp.text[:300]}"
            )
            return resp

        slug = f"uc4b-{uuid4().hex[:8]}"

        # ── Schritt 1: Verein mit billing_notes + IBAN anlegen ──────────────
        org = req(
            "POST", "/admin/organizations", 201, "owner",
            json={
                "slug": slug,
                "name": "Reise-Verein",
                "contact_email": "reise@example.com",
                "iban": "DE89370400440532013000",
                "billing_notes": "Interner Vermerk, nur Plattform-Admin",
            },
        ).json()
        assert "billing_notes" in org, "platform_owner muss billing_notes sehen"

        with cursor() as cur:
            cur.execute(
                "INSERT INTO organization_members(organization_id,user_id) VALUES (%s,%s)",
                (org["id"], admin_a_id),
            )

        req("PATCH", f"/admin/organizations/{slug}", 200, "owner",
            json={"site_published": True})

        # ── Schritt 3: billing_notes nicht in Vereins-Admin-Antwort ─────────
        org_admin_view = req("GET", f"/admin/organizations/{slug}").json()
        assert "billing_notes" not in org_admin_view, (
            "Org-Admin darf billing_notes nicht sehen"
        )

        # ── Schritt 2: Camp (Kapazität 2) anlegen und veröffentlichen ───────
        req(
            "POST", f"/admin/organizations/{slug}/camps", 201,
            json={
                "slug": "sommer",
                "title": "Sommerreise 2027",
                "start_date": "2027-07-05",
                "end_date": "2027-07-09",
                "age_min": 6,
                "age_max": 12,
                "capacity": 2,
                "price_cents": 14900,
            },
        )
        req("PATCH", f"/admin/organizations/{slug}/camps/sommer", 200,
            json={"status": "published"})

        public = f"/api/v1/organizations/{slug}"

        def child_payload(first_name: str, birth: str = "2017-05-10") -> dict:
            return {
                "parent_first_name": "Eva",
                "parent_last_name": "Muster",
                "parent_email": "eva@example.com",
                "parent_phone": "+49 561 123456",
                "child_first_name": first_name,
                "child_last_name": "Muster",
                "child_birth_date": birth,
                "emergency_contact_phone": "+49 561 999888",
                "photo_permission": False,
                "terms_accepted": True,
                "privacy_accepted": True,
            }

        # ── Schritt 4: Kind 1 → registered, CP-Referenz, Mail ohne Namen ────
        r1 = req(
            "POST", f"{public}/camps/sommer/registrations", 201, None,
            json=child_payload("Kind1"),
        ).json()
        assert r1["status"] == "registered"
        expected_ref = f"Sommerreise 2027 CP-{r1['registration_token'][:8].upper()}"
        assert r1["payment_reference"] == expected_ref
        assert "Kind1" not in expected_ref

        # Bestätigungsmail im FakeProvider: kein Kind- oder Elternname im Betreff
        parent_mails = [m for m in fake.sent if m.to_email == "eva@example.com"]
        assert parent_mails, "Keine Elternmail nach Anmeldung"
        subject1 = parent_mails[0].subject
        assert "Kind1" not in subject1, "Kind-Name darf nicht im Betreff stehen"
        assert "Eva" not in subject1, "Elternname darf nicht im Betreff stehen"
        assert "Anmeldebestätigung" in subject1
        assert expected_ref in parent_mails[0].html_content
        assert "Kind1" not in parent_mails[0].html_content

        # Same family, different camp: each response, mail and admin row keeps its own title.
        req(
            "POST", f"/admin/organizations/{slug}/camps", 201,
            json={
                "slug": "winter", "title": "Winterreise 2027",
                "start_date": "2027-01-04", "end_date": "2027-01-07",
                "age_min": 6, "age_max": 12, "capacity": 2, "price_cents": 14900,
            },
        )
        req("PATCH", f"/admin/organizations/{slug}/camps/winter", 200,
            json={"status": "published"})
        winter = req(
            "POST", f"{public}/camps/winter/registrations", 201, None,
            json=child_payload("Kind1"),
        ).json()
        winter_ref = f"Winterreise 2027 CP-{winter['registration_token'][:8].upper()}"
        assert winter["payment_reference"] == winter_ref
        assert winter_ref != expected_ref
        winter_mail = next(m for m in fake.sent if m.to_email == "eva@example.com" and winter_ref in m.html_content)
        assert "Kind1" not in winter_mail.subject + winter_mail.html_content
        assert "Eva" not in winter_mail.subject + winter_mail.html_content
        summer_admin = req("GET", f"/admin/organizations/{slug}/camps/sommer/registrations").json()
        winter_admin = req("GET", f"/admin/organizations/{slug}/camps/winter/registrations").json()
        assert [r["payment_reference"] for r in summer_admin] == [expected_ref]
        assert [r["payment_reference"] for r in winter_admin] == [winter_ref]
        assert all("camp_title" not in r for r in summer_admin + winter_admin)

        # ── Schritt 5: Kind 2 → registered (Camp jetzt voll) ────────────────
        r2 = req(
            "POST", f"{public}/camps/sommer/registrations", 201, None,
            json=child_payload("Kind2"),
        ).json()
        assert r2["status"] == "registered"

        # ── Schritt 6: Kind 3 → waitlist, Wartelistenmail ohne Namen im Betreff
        r3 = req(
            "POST", f"{public}/camps/sommer/registrations", 201, None,
            json=child_payload("Kind3"),
        ).json()
        assert r3["status"] == "waitlist"

        waitlist_mails = [m for m in fake.sent if "Wartelistenbestätigung" in m.subject]
        assert waitlist_mails, "Keine Wartelistenmail nach Wartelisten-Anmeldung"
        assert "Kind3" not in waitlist_mails[0].subject

        # ── Schritt 7: Admin markiert Kind 1 als bezahlt ────────────────────
        tok1 = r1["registration_token"]
        payment_update = req(
            "PATCH",
            f"/admin/organizations/{slug}/camps/sommer/registrations/{tok1}/payment-status",
            200,
            json={"payment_status": "paid"},
        ).json()
        assert payment_update["payment_reference"] == expected_ref
        regs = req("GET", f"/admin/organizations/{slug}/camps/sommer/registrations").json()
        paid_reg = next(r for r in regs if r["registration_token"] == tok1)
        assert paid_reg["payment_status"] == "paid"

        # ── Schritt 8: Admin storniert Kind 2 → Kind 3 rückt nach ──────────
        tok2 = r2["registration_token"]
        tok3 = r3["registration_token"]
        mails_before_cancel = len(fake.sent)
        cancel_result = req(
            "POST",
            f"/admin/organizations/{slug}/camps/sommer/registrations/{tok2}/cancel",
            200,
        ).json()

        assert cancel_result["cancelled"]["status"] == "cancelled"
        # API-Benachrichtigung: promoted-Feld zeigt, dass Kind 3 nachgerückt ist
        assert cancel_result["promoted"] is not None
        assert cancel_result["promoted"]["registration_token"] == tok3
        assert cancel_result["promoted"]["status"] == "registered"

        # Kind 3 ist jetzt registered in der Anmeldungsliste
        regs_after = req(
            "GET", f"/admin/organizations/{slug}/camps/sommer/registrations"
        ).json()
        kind3_row = next(r for r in regs_after if r["registration_token"] == tok3)
        assert kind3_row["status"] == "registered"

        # A2: Storno und Nachrücken informieren beide Familien automatisch
        new_mails = fake.sent[mails_before_cancel:]
        assert [m.subject for m in new_mails] == [
            "Anmeldung storniert – Sommerreise 2027",
            "Platz frei – Ihre Anmeldung ist bestätigt – Sommerreise 2027",
        ]
        assert "Platz frei geworden" in new_mails[1].html_content
        assert kind3_row["email_error"] is None

        # ── Schritt 9: Tagesliste enthält Kind 1 + Kind 3, nicht Kind 2 ─────
        export_resp = req(
            "GET", f"/admin/organizations/{slug}/camps/sommer/export.xlsx?view=daylist"
        )
        wb = load_workbook(BytesIO(export_resp.content))
        sheet = wb.active
        names = {
            sheet.cell(row=r, column=1).value
            for r in range(2, sheet.max_row + 1)
        }
        assert "Kind1 Muster" in names, f"Kind1 fehlt in Tagesliste: {names}"
        assert "Kind3 Muster" in names, f"Kind3 fehlt in Tagesliste: {names}"
        assert not any("Kind2" in (n or "") for n in names), (
            f"Kind2 (storniert) darf nicht in der Tagesliste sein: {names}"
        )

        # ── Schritt 10: Doppelte Anmeldung desselben Kindes → 409 ───────────
        req(
            "POST", f"{public}/camps/sommer/registrations", 409, None,
            json=child_payload("Kind1"),
        )

        # ── Schritt 11: Fremder Vereins-Admin sieht nichts (403) ────────────
        resp_b = client.get(
            f"/admin/organizations/{slug}/camps/sommer/registrations",
            headers=h("admin-b"),
        )
        assert resp_b.status_code in (403, 404), (
            f"Fremder Admin muss 403/404 bekommen, war: {resp_b.status_code}"
        )

        # ── Schritt 12: billing_notes nie in öffentlicher API ───────────────
        pub_org = req("GET", public, identity=None).json()
        assert "billing_notes" not in pub_org
